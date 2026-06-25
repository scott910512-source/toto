import win32com.client
import pythoncom
import os
import re
from datetime import datetime


def sanitize_filename(name):
    return re.sub(r'[\\/*?:"<>|]', "_", name)


def _matches(subject, body, keywords, search_in, keyword_mode):
    # 키워드 없으면 전체 매칭
    if not keywords:
        return True

    targets = []
    if search_in in ("subject", "both"):
        targets.append(subject.lower())
    if search_in in ("body", "both"):
        targets.append(body.lower())
    text = " ".join(targets)

    if keyword_mode == "AND":
        return all(kw.lower() in text for kw in keywords)
    else:  # OR
        return any(kw.lower() in text for kw in keywords)


def _collect_all(folders):
    """하위폴더 포함 전체 수집"""
    result = []
    for folder in folders:
        result.append(folder)
        if folder.Folders.Count > 0:
            result.extend(_collect_all(folder.Folders))
    return result


def _collect_by_name(folders, name_filter):
    """이름 필터에 맞는 폴더만 수집"""
    result = []
    for folder in folders:
        if name_filter.lower() in folder.Name.lower():
            result.append(folder)
        if folder.Folders.Count > 0:
            result.extend(_collect_by_name(folder.Folders, name_filter))
    return result


def _get_target_folders(ns, folder_name_filter=None):
    if not folder_name_filter:
        # 받은편지함 + 모든 하위폴더
        inbox = ns.GetDefaultFolder(6)
        return [inbox] + _collect_all(inbox.Folders)

    # 폴더명 필터: 전체 계정에서 이름 포함 폴더 검색
    all_folders = []
    for store in ns.Stores:
        root = store.GetRootFolder()
        all_folders.extend(_collect_by_name(root.Folders, folder_name_filter))

    if not all_folders:
        raise ValueError(f"'{folder_name_filter}' 이름을 포함한 폴더를 찾을 수 없습니다.")
    return all_folders


def _to_naive_dt(dt):
    """pywintypes.datetime → naive Python datetime 변환"""
    return datetime(dt.year, dt.month, dt.day, dt.hour, dt.minute, dt.second)


def _msg_matches_filters(msg, keywords, search_in, keyword_mode,
                          sender_filter, date_from, date_to,
                          require_attachment):
    subject = msg.Subject or ""
    body = msg.Body or ""
    sender = (msg.SenderEmailAddress or "").lower()
    received = _to_naive_dt(msg.ReceivedTime)

    if date_from and received < date_from:
        return False
    if date_to and received > date_to:
        return False
    if sender_filter and sender_filter.lower() not in sender:
        return False
    if require_attachment and msg.Attachments.Count == 0:
        return False
    if not _matches(subject, body, keywords, search_in, keyword_mode):
        return False
    return True


def preview_emails(keywords, search_in="subject", keyword_mode="OR",
                   folder_name=None, sender_filter=None,
                   date_from=None, date_to=None,
                   require_attachment=False, limit=2):
    pythoncom.CoInitialize()
    try:
        outlook = win32com.client.Dispatch("Outlook.Application")
        ns = outlook.GetNamespace("MAPI")
        folders = _get_target_folders(ns, folder_name)

        previews = []
        for folder in folders:
            messages = folder.Items
            messages.Sort("[ReceivedTime]", True)
            for msg in messages:
                if len(previews) >= limit:
                    break
                try:
                    if _msg_matches_filters(msg, keywords, search_in, keyword_mode,
                                            sender_filter, date_from, date_to,
                                            require_attachment):
                        previews.append({
                            "subject": (msg.Subject or "")[:60],
                            "sender": msg.SenderEmailAddress or "",
                            "received": _to_naive_dt(msg.ReceivedTime).strftime("%Y-%m-%d %H:%M"),
                            "attachments": msg.Attachments.Count,
                            "folder": folder.Name
                        })
                except Exception:
                    continue
            if len(previews) >= limit:
                break
        return previews
    finally:
        pythoncom.CoUninitialize()


def save_attachments(keywords, file_extensions, save_dir,
                     search_in="subject", keyword_mode="OR",
                     folder_name=None, sender_filter=None,
                     date_from=None, date_to=None,
                     require_attachment=False,
                     make_excel=True, dry_run=False,
                     stop_event=None, progress_cb=None, log_cb=None):

    def log(msg, level="INFO"):
        line = f"[{level}] {msg}"
        if log_cb:
            log_cb(line, level)
        else:
            print(line)

    os.makedirs(save_dir, exist_ok=True)
    pythoncom.CoInitialize()
    try:
        outlook = win32com.client.Dispatch("Outlook.Application")
        ns = outlook.GetNamespace("MAPI")
        folders = _get_target_folders(ns, folder_name)

        today_str = datetime.now().strftime("%Y%m%d")
        matched_mails = 0
        saved_files = 0
        skipped_mails = 0
        error_mails = []
        excel_rows = []

        log(f"검색 시작 {'[DRY RUN]' if dry_run else ''}")
        log(f"키워드: {keywords} ({keyword_mode})")
        log(f"저장경로: {save_dir}")
        log("-" * 40)

        total = sum(f.Items.Count for f in folders)
        processed = 0

        for folder in folders:
            log(f"폴더 진입: {folder.Name}")
            messages = folder.Items
            messages.Sort("[ReceivedTime]", True)

            for msg in messages:
                if stop_event and stop_event.is_set():
                    log("사용자 중단 요청으로 종료합니다.", "WARN")
                    break

                processed += 1
                if progress_cb:
                    progress_cb(processed, total)

                try:
                    if not _msg_matches_filters(msg, keywords, search_in, keyword_mode,
                                                sender_filter, date_from, date_to,
                                                require_attachment):
                        continue

                    subject = msg.Subject or ""
                    sender = msg.SenderEmailAddress or ""
                    received = _to_naive_dt(msg.ReceivedTime)
                    received_str = received.strftime("%Y%m%d")
                    safe_subject = sanitize_filename(subject[:40])

                    mail_folder_name = f"{received_str}_{safe_subject}"
                    mail_folder = os.path.join(save_dir, mail_folder_name)

                    txt_name = f"{received_str}_{safe_subject}_{today_str}_본문.txt"
                    txt_path = os.path.join(mail_folder, txt_name)

                    if os.path.exists(mail_folder) and os.path.exists(txt_path):
                        log(f"중복 스킵: {mail_folder_name}", "SKIP")
                        skipped_mails += 1
                        continue

                    matched_mails += 1
                    log(f"[{matched_mails}] {subject[:45]} ({received_str})")

                    if dry_run:
                        log(f"  → DRY RUN: 저장 생략", "INFO")
                        continue

                    os.makedirs(mail_folder, exist_ok=True)

                    with open(txt_path, "w", encoding="utf-8") as f:
                        f.write(f"제목   : {subject}\n")
                        f.write(f"발신자 : {sender}\n")
                        f.write(f"수신일 : {received.strftime('%Y-%m-%d %H:%M:%S')}\n")
                        f.write(f"저장일 : {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
                        f.write("=" * 60 + "\n\n")
                        f.write(msg.Body or "")
                    log(f"  본문: {txt_name}", "INFO")

                    if msg.HTMLBody:
                        html_name = f"{received_str}_{safe_subject}_{today_str}_본문.html"
                        html_path = os.path.join(mail_folder, html_name)
                        with open(html_path, "w", encoding="utf-8") as f:
                            f.write(msg.HTMLBody)

                    att_names = []
                    for att in msg.Attachments:
                        att_name = att.FileName
                        ext = os.path.splitext(att_name)[1].lower()
                        if file_extensions and ext not in file_extensions:
                            log(f"  첨부 스킵(확장자): {att_name}", "SKIP")
                            continue
                        save_name = f"{received_str}_{safe_subject}_{today_str}_{att_name}"
                        save_path = os.path.join(mail_folder, save_name)
                        if os.path.exists(save_path):
                            log(f"  첨부 중복 스킵: {att_name}", "SKIP")
                            continue
                        att.SaveAsFile(save_path)
                        log(f"  첨부: {save_name}", "INFO")
                        saved_files += 1
                        att_names.append(att_name)

                    excel_rows.append({
                        "번호": matched_mails,
                        "수신일": received.strftime("%Y-%m-%d %H:%M"),
                        "제목": subject,
                        "발신자": sender,
                        "폴더": folder.Name,
                        "첨부파일": ", ".join(att_names) if att_names else "없음",
                        "저장경로": mail_folder
                    })

                except Exception as e:
                    err_msg = f"오류: {msg.Subject or '(제목없음)'} → {e}"
                    log(err_msg, "ERROR")
                    error_mails.append(err_msg)
                    continue

            if stop_event and stop_event.is_set():
                break

        if error_mails:
            err_path = os.path.join(save_dir, f"{today_str}_오류목록.txt")
            with open(err_path, "w", encoding="utf-8") as f:
                f.write("\n".join(error_mails))
            log(f"오류 {len(error_mails)}건 → {err_path}", "WARN")

        if make_excel and excel_rows and not dry_run:
            excel_path = _save_excel(excel_rows, save_dir, today_str, keywords)
            log(f"Excel 보고서: {os.path.basename(excel_path)}", "INFO")

        log("=" * 40)
        log(f"완료 | 메일:{matched_mails} 첨부:{saved_files} 스킵:{skipped_mails} 오류:{len(error_mails)}")
        return matched_mails, saved_files, len(error_mails)

    finally:
        pythoncom.CoUninitialize()


def _save_excel(rows, save_dir, today_str, keywords):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        return ""

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "메일 목록"

    headers = ["번호", "수신일", "제목", "발신자", "폴더", "첨부파일", "저장경로"]
    hdr_fill = PatternFill("solid", fgColor="0078D4")
    hdr_font = Font(color="FFFFFF", bold=True, size=10)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    even_fill = PatternFill("solid", fgColor="F0F4FA")

    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = hdr_fill
        c.font = hdr_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    ws.row_dimensions[1].height = 22

    for r, row in enumerate(rows, 2):
        for col, key in enumerate(headers, 1):
            c = ws.cell(row=r, column=col, value=row[key])
            c.alignment = Alignment(vertical="center",
                                    wrap_text=(key in ("제목", "첨부파일")))
            c.border = border
            if r % 2 == 0:
                c.fill = even_fill
        ws.row_dimensions[r].height = 18

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    col_widths = [6, 18, 42, 28, 16, 32, 52]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w

    kw_str = "_".join(keywords)[:20]
    path = os.path.join(save_dir, f"{today_str}_{kw_str}_메일목록.xlsx")
    wb.save(path)
    return path
